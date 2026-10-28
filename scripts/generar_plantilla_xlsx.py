"""Genera plantilla_historial_externo.xlsx (HU-S2-09).

    pip install openpyxl
    python generar_plantilla_xlsx.py plantilla_historial_externo.xlsx
"""
import sys
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

HOJAS = {
    "Estudiante": [
        ("Identificacion", True, "Cedula (10 digitos) o numero de pasaporte (5 a 20 letras o numeros, sin simbolos)."),
        ("TipoIdentificacion", True, "Cedula o Pasaporte."),
        ("Nombres", True, "Nombres del estudiante."),
        ("Apellidos", True, "Apellidos del estudiante."),
        ("InstitucionOrigen", True, "Universidad o institucion de donde viene."),
        ("CarreraOrigen", True, "Carrera que cursaba en la institucion de origen."),
        ("NivelActual", False, "Opcional. Numero entero mayor que 0."),
    ],
    "Asignaturas": [
        ("Codigo", True, "Codigo de la asignatura en la institucion de origen. No se puede repetir."),
        ("Nombre", True, "Nombre de la asignatura."),
        ("Creditos", True, "Numero mayor que 0."),
        ("Horas", True, "Numero entero mayor que 0."),
        ("Nota", True, "Numero entre 0 y 10."),
        ("Periodo", True, "Periodo academico en que la curso. Ejemplo: 2024-2025 CI."),
    ],
    "Contenidos": [
        ("CodigoAsignatura", True, "Debe ser un Codigo que exista en la hoja Asignaturas."),
        ("Tema", True, "Tema o unidad del silabo."),
        ("Subtema", False, "Opcional. Subtema o contenido del tema."),
    ],
}
ANCHOS = {"Identificacion": 18, "TipoIdentificacion": 20, "Nombres": 24, "Apellidos": 24,
          "InstitucionOrigen": 30, "CarreraOrigen": 30, "NivelActual": 13, "Codigo": 18,
          "Nombre": 40, "Creditos": 11, "Horas": 9, "Nota": 9, "Periodo": 18,
          "CodigoAsignatura": 20, "Tema": 40, "Subtema": 50}
FILAS_VALIDADAS = 500


def validacion(ws, col, **kw):
    dv = DataValidation(allow_blank=True, showErrorMessage=True, **kw)
    ws.add_data_validation(dv)
    dv.add(f"{col}2:{col}{FILAS_VALIDADAS}")


def main(salida):
    wb = Workbook()
    wb.remove(wb.active)
    for nombre, columnas in HOJAS.items():
        ws = wb.create_sheet(nombre)
        for i, (titulo, obligatoria, ayuda) in enumerate(columnas, start=1):
            c = ws.cell(row=1, column=i, value=titulo)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F4E78" if obligatoria else "5B7F95")
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.comment = Comment(("Obligatorio. " if obligatoria else "") + ayuda, "Plantilla")
            ws.column_dimensions[c.column_letter].width = ANCHOS.get(titulo, 20)
        ws.freeze_panes = "A2"
    est, asi = wb["Estudiante"], wb["Asignaturas"]
    validacion(est, "B", type="list", formula1='"Cedula,Pasaporte"')
    validacion(est, "G", type="whole", operator="greaterThan", formula1="0")
    validacion(asi, "C", type="decimal", operator="greaterThan", formula1="0")
    validacion(asi, "D", type="whole", operator="greaterThan", formula1="0")
    validacion(asi, "E", type="decimal", operator="between", formula1="0", formula2="10")
    # Identificacion como texto para no perder ceros iniciales
    for fila in range(2, FILAS_VALIDADAS + 1):
        est.cell(row=fila, column=1).number_format = "@"
        asi.cell(row=fila, column=1).number_format = "@"
        wb["Contenidos"].cell(row=fila, column=1).number_format = "@"
    wb.save(salida)


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else "plantilla_historial_externo.xlsx")
